import base64
import io
from datetime import date

from openpyxl import load_workbook

from app.extensions import db
from app.models.course import Course
from app.models.score import Score
from app.services import report_service

from conftest import auth


def _seed_score(app):
    with app.app_context():
        db.session.add(Course(
            course_id='MATH01', course_name='数学', teacher_id='T001', term='2025-2026-2', status=1,
        ))
        db.session.add(Score(
            student_id='S001', course_id='MATH01', score=91,
            exam_date=date(2026, 4, 10), exam_batch='期中', status=1,
        ))
        db.session.commit()


def test_teacher_xlsx_export_is_scoped_and_requires_batch(app, client, login):
    admin = login('admin01')
    teacher = login('teacher01')
    _seed_score(app)
    client.put('/api/admin/teacher-bindings/T001', headers=auth(admin), json={
        'class_names': ['2025级1班'], 'course_ids': ['MATH01'],
    })
    assert client.get('/api/reports/scores/export', headers=auth(teacher)).status_code == 400
    exported = client.get(
        '/api/reports/scores/export?term=2025-2026-2&exam_batch=期中', headers=auth(teacher),
    ).get_json()['data']
    content = base64.b64decode(exported['content_base64'])
    assert content.startswith(b'PK\x03\x04')
    workbook = load_workbook(io.BytesIO(content), read_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    assert exported['row_count'] == 1
    assert rows[1][0] == 'S001'


def test_student_pdf_requires_publication_and_returns_file(app, client, login, monkeypatch):
    admin = login('admin01')
    student = login('student01')
    _seed_score(app)
    before = client.get('/api/reports/students/S001/scores.pdf', headers=auth(student))
    assert before.status_code == 403
    setting = client.post('/api/exam-publish', headers=auth(admin), json={
        'exam_name': '期中发布', 'term': '2025-2026-2', 'exam_batch': '期中',
        'class_name': '2025级1班', 'show_subject_scores': True,
    }).get_json()['data']
    client.post(f"/api/exam-publish/{setting['id']}/publish", headers=auth(admin), json={})
    monkeypatch.setattr(report_service, '_build_pdf', lambda student_row, rows: b'%PDF-test')
    exported = client.get(
        '/api/reports/students/OTHER/scores.pdf?term=2025-2026-2&exam_batch=期中',
        headers=auth(student),
    ).get_json()['data']
    assert base64.b64decode(exported['content_base64']) == b'%PDF-test'
    assert exported['row_count'] == 1


def test_teacher_pdf_excludes_unbound_courses(app, client, login, monkeypatch):
    admin = login('admin01')
    teacher = login('teacher01')
    _seed_score(app)
    with app.app_context():
        db.session.add(Course(
            course_id='ENG01', course_name='英语', teacher_id='T001', term='2025-2026-2', status=1,
        ))
        db.session.add(Score(
            student_id='S001', course_id='ENG01', score=75,
            exam_date=date(2026, 4, 10), exam_batch='期中', status=1,
        ))
        db.session.commit()
    client.put('/api/admin/teacher-bindings/T001', headers=auth(admin), json={
        'class_names': ['2025级1班'], 'course_ids': ['MATH01'],
    })
    monkeypatch.setattr(report_service, '_build_pdf', lambda student_row, rows: b'%PDF-test')
    exported = client.get(
        '/api/reports/students/S001/scores.pdf?term=2025-2026-2&exam_batch=期中',
        headers=auth(teacher),
    ).get_json()['data']
    assert exported['row_count'] == 1


def test_parent_confirmation_counts_and_xlsx_export(app, client, login):
    admin = login('admin01')
    _seed_score(app)
    client.post('/api/parent-bindings', headers=auth(admin), json={
        'parent_user_id': 'P001', 'student_user_id': 'S001',
    })
    setting = client.post('/api/exam-publish', headers=auth(admin), json={
        'exam_name': '期中发布', 'term': '2025-2026-2', 'exam_batch': '期中',
        'class_name': '2025级1班', 'require_parent_signature': True,
    }).get_json()['data']
    client.post(f"/api/exam-publish/{setting['id']}/publish", headers=auth(admin), json={})
    listed = client.get('/api/exam-publish', headers=auth(admin)).get_json()['data']['list'][0]
    assert listed['pending_parent_count'] == 1
    exported = client.get(
        f"/api/exam-publish/{setting['id']}/confirmations/export", headers=auth(admin),
    ).get_json()['data']
    assert exported['pending_count'] == 1
    content = base64.b64decode(exported['content_base64'])
    assert content.startswith(b'PK\x03\x04')
    workbook = load_workbook(io.BytesIO(content), read_only=True)
    assert list(workbook.active.iter_rows(values_only=True))[1][3] == 'S001'
