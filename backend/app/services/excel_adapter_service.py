"""规则识别版智能 Excel 适配器。

本模块只做本地 Excel 分析、识别、转换和报告生成，不写数据库，不调用外部 AI。
"""
from __future__ import annotations

import csv
import re
import zlib
from collections import OrderedDict, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from io import StringIO
from pathlib import Path
from typing import Any


DEFAULT_EXAM_DATE = '2026-06-20'
DEFAULT_EXAM_BATCH = '数据集测试'
DEFAULT_TERM = '2025-2026-2'
DEFAULT_PASSWORD = '123456'
DEFAULT_CLASS_NAME = '未分班'

STUDENT_ALIASES = {
    'student_id': {'学号', '学生编号', '学生id', '学籍号', '编号', 'id', 'stu_id', 'student_id', 'student_no'},
    'real_name': {'姓名', '学生姓名', 'name', 'real_name', 'student_name'},
    'class_name': {'班级', '班别', '行政班', '年级班级', 'class', 'class_name'},
    'grade': {'年级', 'grade'},
    'school_name': {'学校', '校区', 'school', 'school_name'},
    'gender': {'性别', 'gender'},
    'phone': {'手机号', '电话', 'phone', 'mobile'},
    'username': {'用户名', 'username', 'login_name'},
    'password': {'密码', '初始密码', 'password'},
    'seat_no': {'座号'},
    'exam_no': {'考号'},
    'id_card': {'身份证号'},
}

SCORE_ALIASES = {
    'course_id': {'课程编号', '课程号', 'course_id'},
    'course_name': {'课程名称', '科目', '学科', 'subject', 'course_name'},
    'score': {'分数', '成绩', '考试成绩', 'score'},
    'exam_date': {'考试日期', '日期', 'exam_date'},
    'exam_batch': {'考试批次', '批次', '考试名称', '考试类型', 'exam_batch'},
    'term': {'学期', '学年', 'term'},
    'paper_title': {'试卷名称'},
    'full_score': {'满分'},
}

SUBJECT_COLUMNS = OrderedDict([
    ('语文', ('CHN01', '语文')),
    ('数学', ('MATH01', '数学')),
    ('英语', ('ENG01', '英语')),
    ('物理', ('PHY01', '物理')),
    ('化学', ('CHEM01', '化学')),
    ('生物', ('BIO01', '生物')),
    ('历史', ('HIS01', '历史')),
    ('地理', ('GEO01', '地理')),
    ('政治', ('POL01', '政治')),
    ('道德与法治', ('MORAL01', '道德与法治')),
    ('科学', ('SCI01', '科学')),
    ('信息技术', ('IT01', '信息技术')),
    ('计算机', ('CS01', '计算机基础')),
    ('计算机基础', ('CS01', '计算机基础')),
    ('体育', ('PE01', '体育')),
    ('音乐', ('MUS01', '音乐')),
    ('美术', ('ART01', '美术')),
])

UNIVERSITY_KEYWORDS = {
    '大学期末成绩表', '大学物理', '高等数学', '线性代数', '概率论与数理统计',
    '大学英语', '专业导论', '马克思主义基本原理', '中国近现代史纲要',
    '毛泽东思想和中国特色社会主义理论体系概论', '学院', '专业', '学分',
    '绩点', '补考成绩',
}

DERIVED_COLUMNS = {
    '总分', '总成绩', '平均分', '均分', '排名', '名次', '班级排名', '年级排名',
    '等级', '成绩等级', '评价', '评语', '备注', '是否及格', '绩点', '补考成绩',
    '最高分', '最低分', '及格人数', '及格率', '优秀率', '年级平均分', '年级最高分',
    '年级最低分', '年级及格率', '年级优秀率',
}

UNSTORED_FIELDS = {
    'grade': '年级',
    'school_name': '学校/校区',
    'gender': '性别',
    'phone': '手机号/电话',
    'seat_no': '座号',
    'exam_no': '考号',
    'id_card': '身份证号',
    'paper_title': '试卷名称',
    'full_score': '满分',
}


@dataclass
class SheetAnalysis:
    source_file: str
    sheet_name: str
    title_rows: list[list[Any]] = field(default_factory=list)
    header_row_no: int = 0
    columns: list[str] = field(default_factory=list)
    sample_rows: list[list[Any]] = field(default_factory=list)
    sheet_type: str = 'unknown'
    mapping: dict[str, str] = field(default_factory=dict)
    subject_columns: dict[str, dict[str, str]] = field(default_factory=dict)
    derived_columns: list[str] = field(default_factory=list)
    unstored_columns: dict[str, str] = field(default_factory=dict)
    unrecognized_columns: list[str] = field(default_factory=list)
    rows_read: int = 0
    students_count: int = 0
    scores_count: int = 0
    courses_count: int = 0
    warnings: list[str] = field(default_factory=list)


def analyze_excel(file_path: str | Path) -> dict[str, Any]:
    """分析单个 Excel 文件并返回转换结果。"""
    return analyze_workbook(file_path)


def analyze_workbook(file_path: str | Path) -> dict[str, Any]:
    """分析工作簿内所有工作表。"""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError('缺少依赖 openpyxl，请先安装 requirements.txt') from exc

    file_path = Path(file_path)
    workbook = load_workbook(file_path, read_only=True, data_only=True)

    result = _empty_result()
    result['source_files'].append(str(file_path))

    for sheet in workbook.worksheets:
        if hasattr(sheet, 'reset_dimensions'):
            sheet.reset_dimensions()
        sheet_result = analyze_sheet(sheet, file_path.name)
        _merge_result(result, sheet_result)

    return result


def analyze_sheet(sheet, source_file: str) -> dict[str, Any]:
    """分析单个工作表并转换为标准学生、成绩、课程和报告数据。"""
    rows = _read_sheet_rows(sheet)
    rows = _expand_single_column_delimited_rows(rows)
    analysis = _build_sheet_analysis(source_file, sheet.title, rows)

    result = _empty_result()
    result['sheets'].append(_analysis_to_dict(analysis))

    if _is_university_sheet(analysis):
        analysis.sheet_type = 'ignored_non_k12'
        analysis.warnings.append('非 K12 数据已忽略')
        result['sheets'][-1] = _analysis_to_dict(analysis)
        return result

    if not analysis.columns:
        analysis.warnings.append('未识别到有效表头')
        return result

    student_rows = []
    score_rows = []
    course_rows = []
    error_rows = []

    data_rows = rows[analysis.header_row_no:]
    generated_ids: dict[tuple[str, str], str] = {}

    if analysis.sheet_type in {'student_info', 'long_score', 'wide_score', 'mixed'}:
        student_rows, student_errors = convert_students(data_rows, analysis.mapping, generated_ids, analysis)
        error_rows.extend(student_errors)

    if analysis.sheet_type == 'long_score':
        score_rows, course_rows, score_errors = convert_scores_long(data_rows, analysis.mapping, generated_ids, analysis)
        error_rows.extend(score_errors)
    elif analysis.sheet_type in {'wide_score', 'mixed'}:
        score_rows, course_rows, score_errors = convert_scores_wide(data_rows, analysis.mapping, generated_ids, analysis)
        error_rows.extend(score_errors)

    analysis.students_count = len(student_rows)
    analysis.scores_count = len(score_rows)
    analysis.courses_count = len(course_rows)
    result['sheets'][-1] = _analysis_to_dict(analysis)
    result['students'].extend(student_rows)
    result['scores'].extend(score_rows)
    result['courses'].extend(course_rows)
    result['errors'].extend(error_rows)
    return result


def detect_sheet_type(columns: list[str], sample_rows: list[list[Any]] | None = None) -> str:
    """根据字段映射和样例行判断表格类型。"""
    mapping = infer_column_mapping(columns, sample_rows or {})
    subject_cols = _infer_subject_columns(columns)

    has_student_identity = bool({'student_id', 'real_name'} & set(mapping))
    has_student_info = has_student_identity or ('student_id' in mapping and 'class_name' in mapping)
    has_long_score = 'score' in mapping and ('course_id' in mapping or 'course_name' in mapping)
    has_subject_scores = bool(subject_cols)

    if has_long_score and has_student_identity:
        return 'long_score'
    if has_subject_scores and has_student_info:
        return 'mixed'
    if has_subject_scores and has_student_identity:
        return 'wide_score'
    if has_student_info:
        return 'student_info'
    return 'unknown'


def infer_column_mapping(columns: list[str], sample_rows: list[list[Any]] | None = None) -> dict[str, str]:
    """按字段别名词典推断标准字段到原列名的映射。"""
    mapping = {}
    alias_map = {}
    for field_name, aliases in {**STUDENT_ALIASES, **SCORE_ALIASES}.items():
        for alias in aliases:
            alias_map[_normalize_header(alias)] = field_name

    for column in columns:
        normalized = _normalize_header(column)
        field_name = alias_map.get(normalized)
        if field_name and field_name not in mapping:
            mapping[field_name] = column

    return mapping


def convert_students(rows: list[list[Any]], mapping: dict[str, str], generated_ids: dict[tuple[str, str], str],
                     analysis: SheetAnalysis) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """转换学生基础信息。"""
    students = []
    errors = []
    columns = analysis.columns

    for offset, row in enumerate(rows, start=analysis.header_row_no + 1):
        record = _row_to_dict(columns, row)
        if _is_empty_record(record):
            continue

        student_id = _clean(_value_by_field(record, mapping, 'student_id'))
        real_name = _clean(_value_by_field(record, mapping, 'real_name'))
        class_name = _clean(_value_by_field(record, mapping, 'class_name'))

        if not student_id:
            student_id = _stable_student_id(real_name, class_name, generated_ids, offset)
            errors.append(_error(analysis, offset, _column_by_field(mapping, 'student_id'), '', '缺少学号，已生成稳定学号'))

        if not real_name:
            real_name = student_id
            errors.append(_error(analysis, offset, _column_by_field(mapping, 'real_name'), '', '缺少姓名，已用学号作为姓名占位'))

        if not class_name:
            class_name = DEFAULT_CLASS_NAME
            errors.append(_error(analysis, offset, _column_by_field(mapping, 'class_name'), '', '缺少班级，已使用“未分班”占位'))

        username = _clean(_value_by_field(record, mapping, 'username')) or student_id
        password = _clean(_value_by_field(record, mapping, 'password')) or DEFAULT_PASSWORD

        students.append({
            'student_id': student_id,
            'real_name': real_name,
            'class_name': class_name,
            'username': username,
            'password': password,
            '_raw_student_id': student_id,
            '_person_key': _person_key(analysis, student_id, real_name, class_name),
        })

    return students, errors


def convert_scores_long(rows: list[list[Any]], mapping: dict[str, str], generated_ids: dict[tuple[str, str], str],
                        analysis: SheetAnalysis) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    """转换成绩长表。"""
    scores = []
    courses = []
    errors = []
    columns = analysis.columns

    for offset, row in enumerate(rows, start=analysis.header_row_no + 1):
        record = _row_to_dict(columns, row)
        if _is_empty_record(record):
            continue

        student_id = _clean(_value_by_field(record, mapping, 'student_id'))
        real_name = _clean(_value_by_field(record, mapping, 'real_name'))
        class_name = _clean(_value_by_field(record, mapping, 'class_name'))
        if not student_id:
            student_id = _stable_student_id(real_name, class_name, generated_ids, offset)
            errors.append(_error(analysis, offset, _column_by_field(mapping, 'student_id'), '', '缺少学号，成绩行已生成稳定学号'))

        course_name = _clean(_value_by_field(record, mapping, 'course_name'))
        course_id, standard_course_name = _course_id_for_subject(course_name)
        course_id = _clean(_value_by_field(record, mapping, 'course_id')) or course_id
        if not _is_k12_course(course_id, course_name):
            errors.append(_error(analysis, offset, _column_by_field(mapping, 'course_name'), course_name, '非 K12 课程已忽略'))
            continue

        score_value = _value_by_field(record, mapping, 'score')
        score = _parse_score(score_value)

        if score is None:
            errors.append(_error(analysis, offset, _column_by_field(mapping, 'score'), score_value, '分数为空、非数字或不在 0-100 范围'))
            continue

        exam_date = _format_date(_value_by_field(record, mapping, 'exam_date')) or DEFAULT_EXAM_DATE
        exam_batch = _clean(_value_by_field(record, mapping, 'exam_batch')) or DEFAULT_EXAM_BATCH
        term = _clean(_value_by_field(record, mapping, 'term')) or DEFAULT_TERM
        course_name = course_name or standard_course_name or course_id

        scores.append({
            'student_id': student_id,
            'course_id': course_id,
            'score': score,
            'exam_date': exam_date,
            'exam_batch': exam_batch,
            'term': term,
            '_raw_student_id': student_id,
            '_person_key': _person_key(analysis, student_id, real_name, class_name),
        })
        courses.append({'course_id': course_id, 'course_name': course_name, 'term': term})

    return scores, courses, errors


def convert_scores_wide(rows: list[list[Any]], mapping: dict[str, str], generated_ids: dict[tuple[str, str], str],
                        analysis: SheetAnalysis) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    """转换成绩宽表为系统成绩长表。"""
    scores = []
    errors = []
    columns = analysis.columns
    term = _infer_context_value(analysis, 'term') or DEFAULT_TERM
    exam_date = _infer_context_value(analysis, 'exam_date') or DEFAULT_EXAM_DATE
    exam_batch = _infer_context_value(analysis, 'exam_batch') or DEFAULT_EXAM_BATCH

    for offset, row in enumerate(rows, start=analysis.header_row_no + 1):
        record = _row_to_dict(columns, row)
        if _is_empty_record(record):
            continue

        real_name = _clean(_value_by_field(record, mapping, 'real_name'))
        class_name = _clean(_value_by_field(record, mapping, 'class_name'))
        student_id = _clean(_value_by_field(record, mapping, 'student_id'))
        if not student_id:
            student_id = _stable_student_id(real_name, class_name, generated_ids, offset)
            errors.append(_error(analysis, offset, _column_by_field(mapping, 'student_id'), '', '缺少学号，成绩行已生成稳定学号'))

        for column_name, subject_info in analysis.subject_columns.items():
            score_value = record.get(column_name)
            if _clean(score_value) == '':
                continue

            score = _parse_score(score_value)
            if score is None:
                errors.append(_error(analysis, offset, column_name, score_value, '分数为空、非数字或不在 0-100 范围'))
                continue

            scores.append({
                'student_id': student_id,
                'course_id': subject_info['course_id'],
                'score': score,
                'exam_date': exam_date,
                'exam_batch': exam_batch,
                'term': term,
                '_raw_student_id': student_id,
                '_person_key': _person_key(analysis, student_id, real_name, class_name),
            })

    courses = build_course_list(analysis.subject_columns, term)
    return scores, courses, errors


def build_course_list(score_columns: dict[str, dict[str, str]], term: str = DEFAULT_TERM) -> list[dict[str, str]]:
    """从识别到的学科成绩列生成课程清单。"""
    courses = []
    seen = set()
    for info in score_columns.values():
        course_id = info['course_id']
        if course_id in seen:
            continue
        seen.add(course_id)
        courses.append({
            'course_id': course_id,
            'course_name': info['course_name'],
            'term': term,
        })
    return courses


def generate_converted_files(result: dict[str, Any], output_dir: str | Path) -> dict[str, str]:
    """生成标准导入 Excel 和转换报告。"""
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise RuntimeError('缺少依赖 openpyxl，请先安装 requirements.txt') from exc

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    _normalize_student_ids(result)

    students = _dedupe(result['students'], ['student_id'])
    courses = _dedupe(result['courses'], ['course_id'])
    scores = _dedupe(result['scores'], ['student_id', 'course_id', 'exam_batch'])
    errors = result['errors']

    files = {}
    files['students_import_ready'] = str(output_dir / 'students_import_ready.xlsx')
    _write_xlsx(Workbook, files['students_import_ready'], ['student_id', 'real_name', 'class_name', 'username', 'password'], students)

    files['scores_import_ready'] = str(output_dir / 'scores_import_ready.xlsx')
    _write_xlsx(Workbook, files['scores_import_ready'], ['student_id', 'course_id', 'score', 'exam_date', 'exam_batch', 'term'], scores)

    files['courses_required'] = str(output_dir / 'courses_required.xlsx')
    _write_xlsx(Workbook, files['courses_required'], ['course_id', 'course_name', 'term'], courses)

    files['conversion_errors'] = str(output_dir / 'conversion_errors.xlsx')
    _write_xlsx(
        Workbook,
        files['conversion_errors'],
        ['source_file', 'sheet_name', 'row_no', 'column_name', 'raw_value', 'error_reason'],
        errors,
    )

    files['conversion_report'] = write_conversion_report(result, output_dir)
    result['summary'] = {
        'students_count': len(students),
        'classes_count': len({row.get('class_name') for row in students if row.get('class_name')}),
        'courses_count': len(courses),
        'scores_count': len(scores),
        'errors_count': len(errors),
    }
    return files


def write_conversion_report(result: dict[str, Any], output_dir: str | Path) -> str:
    """写入 Markdown 转换报告。"""
    output_dir = Path(output_dir)
    report_name = '真实数据集转换报告_K12.md' if 'k12' in output_dir.name.lower() else '真实数据集转换报告.md'
    report_path = output_dir / report_name

    students = _dedupe(result['students'], ['student_id'])
    courses = _dedupe(result['courses'], ['course_id'])
    scores = _dedupe(result['scores'], ['student_id', 'course_id', 'exam_batch'])
    classes = sorted({row.get('class_name') for row in students if row.get('class_name')})
    ignored = sorted({col for sheet in result['sheets'] for col in sheet.get('derived_columns', [])})
    unstored = sorted({label for sheet in result['sheets'] for label in sheet.get('unstored_columns', {}).values()})
    unrecognized = sorted({col for sheet in result['sheets'] for col in sheet.get('unrecognized_columns', [])})

    lines = [
        '# 真实数据集转换报告',
        '',
        '## 总览',
        '',
        f'- 读取文件数: {len(result["source_files"])}',
        f'- 工作表数: {len(result["sheets"])}',
        f'- 识别学生数: {len(students)}',
        f'- 识别班级数: {len(classes)}',
        f'- 识别课程数: {len(courses)}',
        f'- 生成成绩记录数: {len(scores)}',
        f'- 错误/警告行数: {len(result["errors"])}',
        '',
        '## 读取文件',
        '',
    ]

    lines.extend(f'- {path}' for path in result['source_files'])
    lines.extend(['', '## 工作表分析', ''])

    for sheet in result['sheets']:
        lines.extend([
            f'### {sheet["source_file"]} / {sheet["sheet_name"]}',
            '',
            f'- 判断类型: {sheet["sheet_type"]}',
            f'- 表头行: {sheet["header_row_no"]}',
            f'- 字段列表: {", ".join(sheet["columns"]) if sheet["columns"] else "未识别"}',
            f'- 字段映射: {_format_mapping(sheet["mapping"])}',
            f'- 识别学生数: {sheet["students_count"]}',
            f'- 识别课程数: {sheet["courses_count"]}',
            f'- 生成成绩记录数: {sheet["scores_count"]}',
            f'- 忽略派生列: {", ".join(sheet["derived_columns"]) if sheet["derived_columns"] else "无"}',
            f'- 已识别但暂未入库字段: {", ".join(sheet["unstored_columns"].values()) if sheet["unstored_columns"] else "无"}',
            f'- 无法识别列: {", ".join(sheet["unrecognized_columns"]) if sheet["unrecognized_columns"] else "无"}',
            '',
            '样例行:',
            '',
        ])
        for sample in sheet.get('sample_rows', [])[:3]:
            lines.append(f'- {sample}')
        if sheet.get('warnings'):
            lines.extend(['', '警告:'])
            lines.extend(f'- {warning}' for warning in sheet['warnings'])
        lines.append('')

    lines.extend([
        '## 被忽略的派生列',
        '',
        ', '.join(ignored) if ignored else '无',
        '',
        '## 已识别但暂未入库字段',
        '',
        ', '.join(unstored) if unstored else '无',
        '',
        '## 无法识别列',
        '',
        ', '.join(unrecognized) if unrecognized else '无',
        '',
        '## 错误和警告摘要',
        '',
    ])

    for error in result['errors'][:100]:
        lines.append(
            f'- {error["source_file"]} / {error["sheet_name"]} 第 {error["row_no"]} 行 '
            f'{error["column_name"]}: {error["error_reason"]}，原值={error["raw_value"]}'
        )
    if len(result['errors']) > 100:
        lines.append(f'- 还有 {len(result["errors"]) - 100} 条错误/警告，详见 conversion_errors.xlsx')
    if not result['errors']:
        lines.append('无')

    lines.extend([
        '',
        '## 需要先创建或确认的课程',
        '',
    ])
    if courses:
        lines.extend(f'- {row["course_id"]}: {row["course_name"]} ({row["term"]})' for row in courses)
    else:
        lines.append('未识别到课程。')

    lines.extend([
        '',
        '## 推荐导入顺序',
        '',
        '1. 使用管理员账号导入 `students_import_ready.xlsx`。',
        '2. 导入后确认用户管理中出现学生。',
        '3. 确认成绩查询的班级选择中出现真实班级。',
        '4. 检查 `courses_required.xlsx` 中课程是否已存在。',
        '5. 如果课程不存在，先在管理员课程管理中手动新增课程。',
        '6. 进入“成绩录入 → 批量导入成绩”，导入 `scores_import_ready.xlsx`。',
        '7. 最后测试统计分析、预警名单、AI 学情分析。',
        '',
        '## 说明',
        '',
        '- 本转换器是规则识别版，不调用外部 AI，不写数据库。',
        '- K12 转换范围会整体忽略明显属于大学教务的数据和大学课程。',
        '- 生成的 Excel 只是标准导入模板，仍需用户通过现有导入功能手动确认。',
        '- 总分、平均分、排名、等级、评语等派生列不会写入成绩模板，系统会在导入成绩后自行计算。',
    ])

    report_path.write_text('\n'.join(lines), encoding='utf-8')
    return str(report_path)


def analyze_directory(input_dir: str | Path, output_dir: str | Path | None = None) -> dict[str, Any]:
    """扫描目录下所有 .xlsx 文件并输出转换文件。"""
    input_dir = Path(input_dir)
    output_dir = Path(output_dir) if output_dir else input_dir / 'converted'
    result = _empty_result()

    for file_path in sorted(input_dir.glob('*.xlsx')):
        if file_path.name.startswith('~$'):
            continue
        workbook_result = analyze_workbook(file_path)
        _merge_result(result, workbook_result)

    files = generate_converted_files(result, output_dir)
    result['output_files'] = files
    result['summary'] = {
        'students_count': len(_dedupe(result['students'], ['student_id'])),
        'classes_count': len({row.get('class_name') for row in result['students'] if row.get('class_name')}),
        'courses_count': len(_dedupe(result['courses'], ['course_id'])),
        'scores_count': len(_dedupe(result['scores'], ['student_id', 'course_id', 'exam_batch'])),
        'errors_count': len(result['errors']),
    }
    return result


def _build_sheet_analysis(source_file: str, sheet_name: str, rows: list[list[Any]]) -> SheetAnalysis:
    header_index, columns, title_rows = _detect_header(rows)
    sample_rows = rows[header_index + 1: header_index + 4] if header_index >= 0 else rows[:3]
    mapping = infer_column_mapping(columns, sample_rows)
    subject_columns = _infer_subject_columns(columns)
    derived_columns = [col for col in columns if _normalize_header(col) in {_normalize_header(v) for v in DERIVED_COLUMNS}]
    unstored_columns = {
        field_name: label
        for field_name, label in UNSTORED_FIELDS.items()
        if field_name in mapping
    }
    recognized_columns = set(mapping.values()) | set(subject_columns) | set(derived_columns)
    unrecognized_columns = [col for col in columns if col and col not in recognized_columns]
    sheet_type = detect_sheet_type(columns, sample_rows)

    warnings = []
    if len(columns) <= 1:
        warnings.append('该工作表只识别到一列，可能是源文件列结构异常或数据不完整')

    return SheetAnalysis(
        source_file=source_file,
        sheet_name=sheet_name,
        title_rows=title_rows,
        header_row_no=header_index + 1 if header_index >= 0 else 0,
        columns=columns,
        sample_rows=sample_rows,
        sheet_type=sheet_type,
        mapping=mapping,
        subject_columns=subject_columns,
        derived_columns=derived_columns,
        unstored_columns=unstored_columns,
        unrecognized_columns=unrecognized_columns,
        rows_read=max(0, len(rows) - header_index - 1) if header_index >= 0 else 0,
        warnings=warnings,
    )


def _read_sheet_rows(sheet) -> list[list[Any]]:
    rows = []
    for row in sheet.iter_rows(values_only=True):
        values = list(row)
        while values and _clean(values[-1]) == '':
            values.pop()
        rows.append(values)
    while rows and not rows[-1]:
        rows.pop()
    return rows


def _expand_single_column_delimited_rows(rows: list[list[Any]]) -> list[list[Any]]:
    non_empty = [row for row in rows if row and _clean(row[0]) != '']
    if not non_empty:
        return rows
    max_cells = max(len([cell for cell in row if _clean(cell) != '']) for row in non_empty)
    if max_cells > 1:
        return rows

    delimiters = ['\t', ',', '，', ';', '；', '|']
    best_delimiter = None
    best_score = 1
    for delimiter in delimiters:
        split_lengths = [len(str(row[0]).split(delimiter)) for row in non_empty if delimiter in str(row[0])]
        score = max(split_lengths) if split_lengths else 1
        if score > best_score:
            best_score = score
            best_delimiter = delimiter

    if not best_delimiter:
        return rows

    expanded = []
    for row in rows:
        value = _clean(row[0]) if row else ''
        if not value:
            expanded.append([])
            continue
        if best_delimiter == ',':
            expanded.append(next(csv.reader(StringIO(value))))
        else:
            expanded.append([part.strip() for part in value.split(best_delimiter)])
    return expanded


def _detect_header(rows: list[list[Any]]) -> tuple[int, list[str], list[list[Any]]]:
    best_index = -1
    best_score = 0
    for index, row in enumerate(rows[:30]):
        columns = [_clean(cell) for cell in row]
        score = _header_score(columns)
        if score > best_score:
            best_index = index
            best_score = score

    if best_index < 0:
        return -1, [], rows[:3]

    return best_index, [_clean(cell) for cell in rows[best_index]], rows[:best_index]


def _header_score(columns: list[str]) -> int:
    score = 0
    alias_values = set()
    for aliases in list(STUDENT_ALIASES.values()) + list(SCORE_ALIASES.values()):
        alias_values.update(_normalize_header(alias) for alias in aliases)
    derived = {_normalize_header(item) for item in DERIVED_COLUMNS}
    subjects = {_normalize_header(item) for item in SUBJECT_COLUMNS}

    for column in columns:
        normalized = _normalize_header(column)
        if normalized in alias_values:
            score += 3
        elif normalized in subjects:
            score += 3
        elif normalized in derived:
            score += 1
    return score


def _infer_subject_columns(columns: list[str]) -> dict[str, dict[str, str]]:
    subject_columns = {}
    for column in columns:
        normalized = _normalize_header(column)
        for subject, (course_id, course_name) in SUBJECT_COLUMNS.items():
            subject_key = _normalize_header(subject)
            if normalized == subject_key or normalized == f'{subject_key}成绩' or normalized == f'{subject_key}分数':
                subject_columns[column] = {'course_id': course_id, 'course_name': course_name}
                break
    return subject_columns


def _is_university_sheet(analysis: SheetAnalysis) -> bool:
    text_parts = [analysis.sheet_name]
    text_parts.extend(analysis.columns)
    for sample in analysis.sample_rows:
        text_parts.extend(_clean(cell) for cell in sample)
    text = ' '.join(part for part in text_parts if part)
    return any(keyword in text for keyword in UNIVERSITY_KEYWORDS)


def _is_k12_course(course_id: str, course_name: str) -> bool:
    if not course_id and not course_name:
        return False
    k12_ids = {info[0] for info in SUBJECT_COLUMNS.values()}
    if course_id in k12_ids:
        return True
    normalized = _normalize_header(course_name)
    return any(normalized == _normalize_header(subject) for subject in SUBJECT_COLUMNS)


def _course_id_for_subject(course_name: str) -> tuple[str, str]:
    normalized = _normalize_header(course_name)
    for subject, (course_id, standard_name) in SUBJECT_COLUMNS.items():
        if normalized == _normalize_header(subject):
            return course_id, standard_name
    if not course_name:
        return 'UNKNOWN', '未知课程'
    return f'NONK12{zlib.crc32(course_name.encode("utf-8")) % 100000:05d}', course_name


def _infer_context_value(analysis: SheetAnalysis, field_name: str) -> str:
    for row in analysis.title_rows:
        text = ' '.join(_clean(cell) for cell in row if _clean(cell))
        if not text:
            continue
        if field_name == 'exam_batch' and any(key in text for key in ['期中', '期末', '月考', '考试', '测试']):
            return text
        if field_name == 'term':
            match = re.search(r'20\d{2}[-—]20\d{2}[-—]?[12]?', text)
            if match:
                return match.group(0).replace('—', '-')
        if field_name == 'exam_date':
            match = re.search(r'20\d{2}[-/]\d{1,2}[-/]\d{1,2}', text)
            if match:
                return _format_date(match.group(0)) or ''
    return ''


def _row_to_dict(columns: list[str], row: list[Any]) -> dict[str, Any]:
    return {column: row[index] if index < len(row) else None for index, column in enumerate(columns)}


def _value_by_field(record: dict[str, Any], mapping: dict[str, str], field_name: str) -> Any:
    column = mapping.get(field_name)
    return record.get(column) if column else None


def _column_by_field(mapping: dict[str, str], field_name: str) -> str:
    return mapping.get(field_name, field_name)


def _person_key(analysis: SheetAnalysis, student_id: str, real_name: str, class_name: str) -> str:
    parts = [
        analysis.source_file,
        analysis.sheet_name,
        _clean(student_id),
        _clean(real_name),
        _clean(class_name),
    ]
    return '||'.join(parts)


def _normalize_student_ids(result: dict[str, Any]) -> None:
    """解决跨表相同原始学号对应不同学生的问题，并保持成绩引用一致。"""
    persons = OrderedDict()
    for row in result.get('students', []):
        key = row.get('_person_key') or _person_key(
            SheetAnalysis(row.get('_source_file', ''), row.get('_sheet_name', '')),
            row.get('student_id', ''),
            row.get('real_name', ''),
            row.get('class_name', ''),
        )
        if key not in persons:
            persons[key] = {
                'raw_student_id': _clean(row.get('_raw_student_id') or row.get('student_id')),
                'real_name': _clean(row.get('real_name')),
                'class_name': _clean(row.get('class_name')),
            }

    raw_to_keys = defaultdict(list)
    for key, info in persons.items():
        raw_to_keys[info['raw_student_id']].append(key)

    id_map = {}
    generated_seq = 1
    used_ids = set()
    for raw_id, keys in raw_to_keys.items():
        if raw_id and len(keys) == 1 and raw_id not in used_ids:
            id_map[keys[0]] = raw_id
            used_ids.add(raw_id)
            continue

        for key in keys:
            while True:
                generated_id = f'K12S{generated_seq:06d}'
                generated_seq += 1
                if generated_id not in used_ids:
                    break
            id_map[key] = generated_id
            used_ids.add(generated_id)
            if raw_id:
                info = persons[key]
                result['errors'].append({
                    'source_file': '智能Excel适配器',
                    'sheet_name': '全局学生ID规范化',
                    'row_no': 0,
                    'column_name': 'student_id',
                    'raw_value': raw_id,
                    'error_reason': f'原始学号在多个不同学生中重复，已为 {info["real_name"]}/{info["class_name"]} 生成 {generated_id}',
                })

    for row in result.get('students', []):
        key = row.get('_person_key')
        normalized_id = id_map.get(key, row.get('student_id'))
        row['student_id'] = normalized_id
        row['username'] = normalized_id

    for row in result.get('scores', []):
        key = row.get('_person_key')
        row['student_id'] = id_map.get(key, row.get('student_id'))


def _stable_student_id(real_name: str, class_name: str, generated_ids: dict[tuple[str, str], str], row_no: int) -> str:
    key = (_clean(real_name), _clean(class_name))
    if key[0] or key[1]:
        if key not in generated_ids:
            generated_ids[key] = f'S{len(generated_ids) + 1:04d}'
        return generated_ids[key]
    return f'S{row_no:04d}'


def _parse_score(value: Any) -> float | None:
    if value is None or _clean(value) == '':
        return None
    try:
        score = float(str(value).strip())
    except (TypeError, ValueError):
        return None
    if score < 0 or score > 100:
        return None
    return round(score, 2)


def _format_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d')
    if isinstance(value, date):
        return value.strftime('%Y-%m-%d')
    text = _clean(value)
    if not text:
        return ''
    for fmt in ['%Y-%m-%d', '%Y/%m/%d', '%Y.%m.%d']:
        try:
            return datetime.strptime(text, fmt).strftime('%Y-%m-%d')
        except ValueError:
            pass
    return ''


def _clean(value: Any) -> str:
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalize_header(value: Any) -> str:
    text = _clean(value).lower()
    return re.sub(r'[\s_（）()【】\[\]:：/\\-]+', '', text)


def _is_empty_record(record: dict[str, Any]) -> bool:
    return all(_clean(value) == '' for value in record.values())


def _error(analysis: SheetAnalysis, row_no: int, column_name: str, raw_value: Any, error_reason: str) -> dict[str, Any]:
    return {
        'source_file': analysis.source_file,
        'sheet_name': analysis.sheet_name,
        'row_no': row_no,
        'column_name': column_name,
        'raw_value': _clean(raw_value),
        'error_reason': error_reason,
    }


def _empty_result() -> dict[str, Any]:
    return {
        'source_files': [],
        'sheets': [],
        'students': [],
        'scores': [],
        'courses': [],
        'errors': [],
        'output_files': {},
        'summary': {},
    }


def _merge_result(target: dict[str, Any], source: dict[str, Any]) -> None:
    for key in ['source_files', 'sheets', 'students', 'scores', 'courses', 'errors']:
        target[key].extend(source.get(key, []))


def _analysis_to_dict(analysis: SheetAnalysis) -> dict[str, Any]:
    return {
        'source_file': analysis.source_file,
        'sheet_name': analysis.sheet_name,
        'header_row_no': analysis.header_row_no,
        'columns': analysis.columns,
        'sample_rows': analysis.sample_rows,
        'sheet_type': analysis.sheet_type,
        'mapping': analysis.mapping,
        'subject_columns': analysis.subject_columns,
        'derived_columns': analysis.derived_columns,
        'unstored_columns': analysis.unstored_columns,
        'unrecognized_columns': analysis.unrecognized_columns,
        'rows_read': analysis.rows_read,
        'students_count': analysis.students_count,
        'scores_count': analysis.scores_count,
        'courses_count': analysis.courses_count,
        'warnings': analysis.warnings,
    }


def _dedupe(rows: list[dict[str, Any]], key_fields: list[str]) -> list[dict[str, Any]]:
    deduped = OrderedDict()
    for row in rows:
        key = tuple(row.get(field) for field in key_fields)
        if not all(key):
            continue
        if key not in deduped:
            deduped[key] = row
    return list(deduped.values())


def _write_xlsx(workbook_class, path: str, headers: list[str], rows: list[dict[str, Any]]) -> None:
    workbook = workbook_class()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, '') for header in headers])
    workbook.save(path)


def _format_mapping(mapping: dict[str, str]) -> str:
    if not mapping:
        return '无'
    return ', '.join(f'{field} <- {column}' for field, column in mapping.items())
