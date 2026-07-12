"""Excel 学生基础信息批量导入服务。"""
import json
import logging
import os
import time
import uuid

from app.extensions import db
from app.models.role import Role, UserRole
from app.models.user import User
from app.services import audit_service, import_batch_service
from app.utils.errors import BusinessError, ErrorCode
from app.utils.hash_util import hash_password
from app.utils.upload_security import validate_xlsx_upload
from app.utils import runtime_storage

logger = logging.getLogger(__name__)

HEADER_ALIASES = {
    'student_id': {'学号', 'student_id'},
    'real_name': {'姓名', 'real_name', 'student_name'},
    'class_name': {'班级', 'class_name'},
    'username': {'用户名', 'username'},
    'grade': {'年级', 'grade'},
    'school_name': {'学校', 'school_name'},
    'gender': {'性别', 'gender'},
    'phone': {'手机号', 'phone'},
    'password': {'初始密码', 'password'},
}
REQUIRED_FIELDS = ['student_id', 'real_name', 'class_name']
DEFAULT_PASSWORD = '123456'


def preview_import(file_storage, operator_id):
    """解析并校验 .xlsx 学生基础信息文件，不写入 users 表。"""
    _validate_file(file_storage)

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise BusinessError(ErrorCode.INTERNAL_SERVER_ERROR, '缺少依赖 openpyxl，请先安装后端依赖') from exc

    try:
        workbook = load_workbook(file_storage.stream, read_only=True, data_only=True)
        sheet = workbook.active
    except Exception as exc:
        raise BusinessError(ErrorCode.USER_IMPORT_INVALID, 'Excel 文件解析失败，请确认文件格式为 .xlsx') from exc

    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        raise BusinessError(ErrorCode.USER_IMPORT_INVALID, 'Excel 文件为空')

    header_map = _build_header_map(header_row)
    missing = [field for field in REQUIRED_FIELDS if field not in header_map]
    if missing:
        raise BusinessError(
            ErrorCode.USER_IMPORT_INVALID,
            f'缺少必填表头: {", ".join(missing)}'
        )

    rows = []
    seen_student_ids = set()
    seen_usernames = set()
    total_rows = 0
    valid_rows = 0
    duplicate_rows = 0
    error_rows = 0

    for row_index, row_values in enumerate(rows_iter, start=2):
        if _is_empty_row(row_values):
            continue

        total_rows += 1
        row = _parse_row(row_index, row_values, header_map)
        _validate_row(row, seen_student_ids, seen_usernames)

        if row['status'] == 'valid':
            valid_rows += 1
        elif row['status'] == 'duplicate':
            duplicate_rows += 1
        else:
            error_rows += 1

        rows.append(row)

    for row in rows:
        password = row.pop('password', None) or DEFAULT_PASSWORD
        if row.get('status') == 'valid':
            row['password_hash'] = hash_password(password)

    import_id = uuid.uuid4().hex
    result = {
        'import_id': import_id,
        'source_filename': os.path.basename(str(file_storage.filename or 'students.xlsx')),
        'total_rows': total_rows,
        'valid_rows': valid_rows,
        'duplicate_rows': duplicate_rows,
        'error_rows': error_rows,
        'rows': rows,
    }
    cached_result = dict(result)
    cached_result['operator_id'] = operator_id
    cached_result['created_at'] = int(time.time())
    _save_preview(import_id, cached_result)
    return _public_preview(result)


def confirm_import(import_id, operator_id, trace_id):
    """导入 preview 中合法且数据库未重复的学生账号。"""
    preview = _load_preview(import_id, operator_id)
    rows = preview.get('rows') or []
    student_role = Role.query.filter_by(role_name='student').first()
    if not student_role:
        raise BusinessError(ErrorCode.ROLE_MAPPING_INVALID, '学生角色不存在，无法导入')

    inserted_count = 0
    skipped_count = 0
    batch = None
    try:
        batch = import_batch_service.create_batch(
            import_type='users',
            operator_user_id=operator_id,
            source_filename=preview.get('source_filename'),
            total_rows=preview.get('total_rows', len(rows)),
            metadata={
                'preview_import_id': import_id,
                'duplicate_rows': preview.get('duplicate_rows', 0),
            },
        )
        for row in rows:
            if row.get('status') != 'valid':
                if row.get('status') == 'duplicate':
                    skipped_count += 1
                continue
            student_id = row['student_id']
            username = row['username']
            if User.query.filter_by(user_id=student_id).first() or User.query.filter_by(username=username).first():
                skipped_count += 1
                continue
            user = User(
                user_id=student_id,
                username=username,
                password_hash=row.get('password_hash') or hash_password(DEFAULT_PASSWORD),
                real_name=row['real_name'],
                class_name=row['class_name'],
                status=1,
            )
            db.session.add(user)
            db.session.flush()
            db.session.add(UserRole(user_id=student_id, role_id=student_role.role_id))
            db.session.flush()
            import_batch_service.record_user_entry(batch, user, row_number=row.get('row_no'))
            inserted_count += 1

        failed_count = max(int(preview.get('total_rows', len(rows))) - inserted_count, 0)
        import_batch_service.complete_batch(batch, inserted_count, failed_count)
        audit_service.write(
            action='user.import.confirm',
            operator_id=operator_id,
            target_type='import_batch',
            target_id=batch.import_batch_id,
            detail={
                'preview_import_id': import_id,
                'inserted_count': inserted_count,
                'skipped_count': skipped_count,
                'error_count': preview.get('error_rows', 0),
            },
            trace_id=trace_id,
        )
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise
    _delete_preview(import_id)

    return {
        'import_id': import_id,
        'import_batch_id': batch.import_batch_id,
        'inserted_count': inserted_count,
        'skipped_count': skipped_count,
        'error_count': preview.get('error_rows', 0),
        'task_status': 'done',
    }


def _validate_file(file_storage):
    try:
        validate_xlsx_upload(file_storage, max_size_mb=10)
    except BusinessError as exc:
        raise BusinessError(ErrorCode.USER_IMPORT_INVALID, exc.message, http_status=exc.http_status) from exc


def _build_header_map(header_row):
    header_map = {}
    alias_to_field = {}
    for field, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            alias_to_field[_normalize_header(alias)] = field

    for index, cell_value in enumerate(header_row):
        field = alias_to_field.get(_normalize_header(cell_value))
        if field and field not in header_map:
            header_map[field] = index

    return header_map


def _parse_row(row_no, row_values, header_map):
    row = {
        'row_no': row_no,
        'student_id': _get_cell(row_values, header_map, 'student_id'),
        'username': _get_cell(row_values, header_map, 'username'),
        'real_name': _get_cell(row_values, header_map, 'real_name'),
        'class_name': _get_cell(row_values, header_map, 'class_name'),
        'grade': _get_cell(row_values, header_map, 'grade'),
        'school_name': _get_cell(row_values, header_map, 'school_name'),
        'gender': _get_cell(row_values, header_map, 'gender'),
        'phone': _get_cell(row_values, header_map, 'phone'),
        'password': _get_cell(row_values, header_map, 'password'),
        'role': 'student',
        'status_value': 1,
        'status': 'valid',
        'errors': [],
    }

    for key in ['student_id', 'username', 'real_name', 'class_name', 'grade',
                'school_name', 'gender', 'phone', 'password']:
        if row[key] is not None:
            row[key] = str(row[key]).strip()

    if not row['username'] and row['student_id']:
        row['username'] = row['student_id']
    if not row['password']:
        row['password'] = DEFAULT_PASSWORD

    return row


def _validate_row(row, seen_student_ids, seen_usernames):
    errors = []
    student_id = row.get('student_id') or ''
    username = row.get('username') or ''

    if not student_id:
        errors.append('学号不能为空')
    if not row.get('real_name'):
        errors.append('姓名不能为空')
    if not row.get('class_name'):
        errors.append('班级不能为空')
    if not username:
        errors.append('用户名不能为空')

    if student_id:
        if student_id in seen_student_ids:
            errors.append('Excel 文件内学号重复')
        else:
            seen_student_ids.add(student_id)

    if username:
        if username in seen_usernames:
            errors.append('Excel 文件内用户名重复')
        else:
            seen_usernames.add(username)

    if errors:
        row['status'] = 'error'
        row['errors'] = errors
        return

    if User.query.filter_by(user_id=student_id).first():
        row['status'] = 'duplicate'
        row['errors'] = ['学号已存在']
        return

    username_owner = User.query.filter_by(username=username).first()
    if username_owner:
        row['status'] = 'error'
        row['errors'] = ['用户名已存在']
        return

    row['status'] = 'valid'
    row['errors'] = []


def _get_cell(row_values, header_map, field):
    index = header_map.get(field)
    if index is None or index >= len(row_values):
        return None
    return row_values[index]


def _normalize_header(value):
    return str(value or '').strip().lower()


def _is_empty_row(row_values):
    return all(value is None or str(value).strip() == '' for value in row_values)


def _preview_dir():
    return runtime_storage.resolve_private_dir('USER_IMPORT_PREVIEW_DIR', 'user_imports')


def _preview_path(import_id):
    return runtime_storage.preview_path(
        'user_imports', import_id, 'USER_IMPORT_PREVIEW_DIR', ErrorCode.USER_IMPORT_INVALID,
    )


def _save_preview(import_id, data):
    runtime_storage.save_preview(
        'user_imports', import_id, data, 'USER_IMPORT_PREVIEW_DIR', ErrorCode.USER_IMPORT_INVALID,
    )


def _load_preview(import_id, operator_id):
    if not import_id:
        raise BusinessError(ErrorCode.USER_IMPORT_INVALID, 'import_id 不能为空')

    return runtime_storage.load_preview(
        'user_imports', import_id, operator_id, 'USER_IMPORT_PREVIEW_DIR',
        ErrorCode.USER_IMPORT_INVALID, '导入预览结果不存在或已过期',
    )


def _delete_preview(import_id):
    if not runtime_storage.delete_preview(
        'user_imports', import_id, 'USER_IMPORT_PREVIEW_DIR', ErrorCode.USER_IMPORT_INVALID,
    ):
        logger.warning('学生导入预览临时文件清理失败: import_id=%s', import_id)


def _public_preview(data):
    public = dict(data)
    public['rows'] = []
    for row in data.get('rows') or []:
        public_row = dict(row)
        public_row.pop('password', None)
        public_row.pop('password_hash', None)
        public['rows'].append(public_row)
    return public
