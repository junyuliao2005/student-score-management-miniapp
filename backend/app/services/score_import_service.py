"""Excel 成绩批量导入服务。"""
import json
import logging
import os
import time
import uuid
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from app.extensions import db
from app.models.course import Course
from app.models.score import Score
from app.models.user import User
from app.services import audit_service, score_service, teacher_scope_service, import_batch_service
from app.utils.errors import BusinessError, ErrorCode
from app.utils.upload_security import validate_xlsx_upload
from app.utils import runtime_storage

logger = logging.getLogger(__name__)

HEADER_ALIASES = {
    'student_id': {'学号', 'student_id'},
    'student_name': {'姓名', 'student_name'},
    'course_id': {'课程编号', 'course_id'},
    'course_name': {'课程名称', 'course_name'},
    'score': {'分数', 'score'},
    'exam_date': {'考试日期', 'exam_date'},
    'exam_batch': {'考试批次', 'exam_batch'},
    'term': {'学期', 'term'},
}
REQUIRED_FIELDS = ['student_id', 'course_id', 'score', 'exam_date', 'exam_batch']


def preview_import(file_storage, current_user):
    """解析并校验 .xlsx 文件，不写入成绩表。"""
    _validate_file(file_storage)

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise BusinessError(ErrorCode.INTERNAL_SERVER_ERROR, '缺少依赖 openpyxl，请先安装后端依赖') from exc

    try:
        workbook = load_workbook(file_storage.stream, read_only=True, data_only=True)
        sheet = workbook.active
    except Exception as exc:
        raise BusinessError(ErrorCode.SCORE_IMPORT_INVALID, 'Excel 文件解析失败，请确认文件格式为 .xlsx') from exc

    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        raise BusinessError(ErrorCode.SCORE_IMPORT_INVALID, 'Excel 文件为空')

    header_map = _build_header_map(header_row)
    missing = [field for field in REQUIRED_FIELDS if field not in header_map]
    if missing:
        raise BusinessError(
            ErrorCode.SCORE_IMPORT_INVALID,
            f'缺少必填表头: {", ".join(missing)}'
        )

    rows = []
    seen_keys = set()
    total_rows = 0
    valid_rows = 0
    error_rows = 0
    duplicate_rows = 0

    for row_index, row_values in enumerate(rows_iter, start=2):
        if _is_empty_row(row_values):
            continue

        total_rows += 1
        row = _parse_row(row_index, row_values, header_map)
        _validate_row(row, seen_keys, current_user)

        if row['status'] == 'valid':
            valid_rows += 1
        elif row['status'] == 'duplicate':
            duplicate_rows += 1
        else:
            error_rows += 1

        rows.append(row)

    import_id = uuid.uuid4().hex
    result = {
        'import_id': import_id,
        'source_filename': os.path.basename(str(file_storage.filename or 'scores.xlsx')),
        'total_rows': total_rows,
        'valid_rows': valid_rows,
        'error_rows': error_rows,
        'duplicate_rows': duplicate_rows,
        'rows': rows,
    }
    cached_result = dict(result)
    cached_result['operator_id'] = current_user.get('user_id')
    cached_result['created_at'] = int(time.time())
    _save_preview(import_id, cached_result)
    return result


def confirm_import(import_id, operator_id, trace_id, current_user):
    """导入 preview 中合法且不重复的成绩。"""
    preview = _load_preview(import_id, operator_id)
    rows = preview.get('rows') or []

    inserted_count = 0
    skipped_count = 0
    original_error_count = int(preview.get('error_rows', 0) or 0)
    batch = None
    try:
        batch = import_batch_service.create_batch(
            import_type='scores',
            operator_user_id=operator_id,
            source_filename=preview.get('source_filename'),
            total_rows=preview.get('total_rows', len(rows)),
            metadata={
                'preview_import_id': import_id,
                'duplicate_rows': preview.get('duplicate_rows', 0),
            },
        )
        for row in rows:
            if row.get('status') != 'valid':
                if row.get('status') == 'duplicate':
                    skipped_count += 1
                continue
            if _has_active_duplicate(row['student_id'], row['course_id'], row['exam_batch']):
                skipped_count += 1
                continue

            result = score_service.create_score(
                {
                    'student_id': row['student_id'],
                    'course_id': row['course_id'],
                    'score': row['score'],
                    'exam_date': row['exam_date'],
                    'exam_batch': row['exam_batch'],
                },
                operator_id=operator_id,
                trace_id=trace_id,
                current_user=current_user,
                commit=False,
            )
            record = Score.query.filter_by(score_id=result['score_id']).first()
            import_batch_service.record_score_entry(batch, record, row_number=row.get('row_no'))
            inserted_count += 1

        failed_count = max(int(preview.get('total_rows', len(rows))) - inserted_count, 0)
        import_batch_service.complete_batch(batch, inserted_count, failed_count)
        audit_service.write(
            action='score.import.confirm',
            operator_id=operator_id,
            target_type='import_batch',
            target_id=batch.import_batch_id,
            detail={
                'preview_import_id': import_id,
                'inserted_count': inserted_count,
                'skipped_count': skipped_count,
                'error_count': original_error_count,
            },
            trace_id=trace_id,
        )
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise
    _delete_preview(import_id)

    return {
        'import_id': import_id,
        'import_batch_id': batch.import_batch_id,
        'inserted_count': inserted_count,
        'skipped_count': skipped_count,
        'error_count': original_error_count,
        'task_status': 'done',
    }


def _validate_file(file_storage):
    try:
        validate_xlsx_upload(file_storage, max_size_mb=10)
    except BusinessError as exc:
        raise BusinessError(ErrorCode.SCORE_IMPORT_INVALID, exc.message, http_status=exc.http_status) from exc


def _build_header_map(header_row):
    header_map = {}
    alias_to_field = {}
    for field, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            alias_to_field[_normalize_header(alias)] = field

    for index, cell_value in enumerate(header_row):
        normalized = _normalize_header(cell_value)
        field = alias_to_field.get(normalized)
        if field and field not in header_map:
            header_map[field] = index

    return header_map


def _parse_row(row_no, row_values, header_map):
    row = {
        'row_no': row_no,
        'student_id': _get_cell(row_values, header_map, 'student_id'),
        'student_name': _get_cell(row_values, header_map, 'student_name'),
        'course_id': _get_cell(row_values, header_map, 'course_id'),
        'course_name': _get_cell(row_values, header_map, 'course_name'),
        'score': _get_cell(row_values, header_map, 'score'),
        'exam_date': _get_cell(row_values, header_map, 'exam_date'),
        'exam_batch': _get_cell(row_values, header_map, 'exam_batch'),
        'term': _get_cell(row_values, header_map, 'term'),
        'status': 'valid',
        'errors': [],
    }

    for key in ['student_id', 'student_name', 'course_id', 'course_name', 'exam_batch', 'term']:
        if row[key] is not None:
            row[key] = str(row[key]).strip()

    return row


def _validate_row(row, seen_keys, current_user):
    errors = []

    student_id = row.get('student_id') or ''
    course_id = row.get('course_id') or ''
    exam_batch = row.get('exam_batch') or ''

    if not student_id:
        errors.append('学号不能为空')
    if not course_id:
        errors.append('课程编号不能为空')
    if not exam_batch:
        errors.append('考试批次不能为空')

    score_value = _parse_score(row.get('score'), errors)
    if score_value is not None:
        row['score'] = float(score_value)

    exam_date = _parse_exam_date(row.get('exam_date'), errors)
    if exam_date is not None:
        row['exam_date'] = exam_date

    student = None
    if student_id:
        student = User.query.filter_by(user_id=student_id).first()
        if not student:
            errors.append('学号不存在')
        elif student.status != 1:
            errors.append('学生状态不是启用')
        elif not student.has_role('student'):
            errors.append('用户不是学生角色')
        elif row.get('student_name') and row['student_name'] != student.real_name:
            errors.append('姓名与学号不匹配')

    course = None
    if course_id:
        course = Course.query.filter_by(course_id=course_id).first()
        if not course:
            errors.append('课程编号不存在')
        elif course.status != 1:
            errors.append('课程已停用')
        else:
            if row.get('course_name') and row['course_name'] != course.course_name:
                errors.append('课程名称与课程编号不匹配')
            if row.get('term') and row['term'] != course.term:
                errors.append('学期与课程所属学期不匹配')

    if student and course and not errors:
        try:
            teacher_scope_service.ensure_access(
                current_user, student_id=student_id, course_id=course_id,
            )
        except BusinessError as exc:
            errors.append(exc.message)

    key = (student_id, course_id, exam_batch)
    if all(key):
        if key in seen_keys:
            errors.append('Excel 文件内存在重复成绩')
        else:
            seen_keys.add(key)

    if errors:
        row['status'] = 'error'
        row['errors'] = errors
        return

    if _has_active_duplicate(student_id, course_id, exam_batch):
        row['status'] = 'duplicate'
        row['errors'] = ['数据库中已存在同一学生、课程、批次的有效成绩']
        return

    row['status'] = 'valid'
    row['errors'] = []


def _parse_score(value, errors):
    try:
        score = Decimal(str(value).strip())
    except (InvalidOperation, ValueError, TypeError, AttributeError):
        errors.append('分数必须是数字')
        return None

    if score < 0 or score > 100:
        errors.append('分数必须在 0 到 100 之间')
        return None

    return score


def _parse_exam_date(value, errors):
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d')
    if isinstance(value, date):
        return value.strftime('%Y-%m-%d')

    try:
        parsed = datetime.strptime(str(value).strip(), '%Y-%m-%d')
        return parsed.strftime('%Y-%m-%d')
    except (ValueError, TypeError, AttributeError):
        errors.append('考试日期必须为 YYYY-MM-DD')
        return None


def _has_active_duplicate(student_id, course_id, exam_batch):
    return Score.query.filter_by(
        student_id=student_id,
        course_id=course_id,
        exam_batch=exam_batch,
        status=1,
    ).first() is not None


def _get_cell(row_values, header_map, field):
    index = header_map.get(field)
    if index is None or index >= len(row_values):
        return None
    return row_values[index]


def _normalize_header(value):
    return str(value or '').strip().lower()


def _is_empty_row(row_values):
    return all(value is None or str(value).strip() == '' for value in row_values)


def _preview_dir():
    return runtime_storage.resolve_private_dir('SCORE_IMPORT_PREVIEW_DIR', 'score_imports')


def _preview_path(import_id):
    return runtime_storage.preview_path(
        'score_imports', import_id, 'SCORE_IMPORT_PREVIEW_DIR', ErrorCode.SCORE_IMPORT_INVALID,
    )


def _save_preview(import_id, data):
    runtime_storage.save_preview(
        'score_imports', import_id, data, 'SCORE_IMPORT_PREVIEW_DIR', ErrorCode.SCORE_IMPORT_INVALID,
    )


def _load_preview(import_id, operator_id):
    if not import_id:
        raise BusinessError(ErrorCode.SCORE_IMPORT_INVALID, 'import_id 不能为空')

    return runtime_storage.load_preview(
        'score_imports', import_id, operator_id, 'SCORE_IMPORT_PREVIEW_DIR',
        ErrorCode.SCORE_IMPORT_INVALID, '导入预览结果不存在或已过期',
    )


def _delete_preview(import_id):
    if not runtime_storage.delete_preview(
        'score_imports', import_id, 'SCORE_IMPORT_PREVIEW_DIR', ErrorCode.SCORE_IMPORT_INVALID,
    ):
        logger.warning('成绩导入预览临时文件清理失败: import_id=%s', import_id)
